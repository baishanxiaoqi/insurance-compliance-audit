"""根据审查点种子表构造覆盖型 smoke 评测数据。"""

from __future__ import annotations

import argparse
import json
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
import sys

from openpyxl import load_workbook

if __package__ in {None, ""}:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from benchmark.schemas import CaseRecord


@dataclass(frozen=True)
class PointSeed:
    category: str
    violation_phrase: str
    violation_text: str
    claim_type: str
    actor_scope: str = "agent"
    compliant_text: str = ""
    keywords: tuple[str, ...] = ()


POINT_SEEDS: dict[str, PointSeed] = {
    "1.1.1": PointSeed("financial_product_confusion", "每年存2万", "这份年金险每年存2万，连续存10年，后面到时间就能慢慢领。", "payment_confusion", compliant_text="这份年金保险年交保费2万元，连续交10年，保障责任和利益以合同为准。", keywords=("存2万", "交保费")),
    "1.1.2": PointSeed("financial_product_confusion", "本金放进去以后按复利慢慢长", "这笔本金放进去以后按复利慢慢长，后面还能拿利息。", "benefit_confusion", compliant_text="保单利益和现金价值不是存款利息，具体以合同条款和利益演示为准。", keywords=("本金", "复利", "利息")),
    "1.1.3": PointSeed("financial_product_confusion", "储蓄账户", "您先把钱放进这个储蓄账户，后面都在账户里滚存。", "account_confusion", compliant_text="万能账户属于保险合同项下账户，不是银行储蓄账户，结算规则以条款披露为准。", keywords=("储蓄账户", "账户")),
    "1.2.1": PointSeed("financial_product_confusion", "就像银行理财，比基金更稳", "这款保险就像银行理财，比基金更稳，拿着更放心。", "product_comparison", compliant_text="这是保险产品，不等同于银行理财或基金，保障责任和收益特征都有区别。", keywords=("银行理财", "基金")),
    "2.1.1": PointSeed("responsibility_exaggeration", "什么情况都能赔", "这个产品基本什么情况都能赔，理赔范围特别宽。", "coverage_exaggeration", compliant_text="是否赔付要看保险责任、免责和理赔条件，不能简单理解为什么情况都能赔。", keywords=("都能赔", "没有限制")),
    "2.2.1": PointSeed("responsibility_exaggeration", "债务问题也能一起规划", "这份保单配上以后，债务问题也能一起规划。", "debt_planning_claim", compliant_text="保险保障不等于债务规划，相关责任边界需要按合同理解。", keywords=("债务规划",)),
    "2.2.2": PointSeed("responsibility_exaggeration", "把税务提前规划好", "买这份保险，相当于顺手把税务也提前规划好了。", "tax_planning_claim", compliant_text="保险产品不能直接宣传为税务规划工具，涉税问题需按政策和专业意见处理。", keywords=("税务规划",)),
    "2.2.3": PointSeed("responsibility_exaggeration", "什么风险都能保住", "有了这张保单，今后大大小小的风险基本都能保住。", "all_risk_coverage_claim", compliant_text="保险只能保障合同约定责任，不存在什么风险都能覆盖的说法。", keywords=("什么风险都能保住",)),
    "2.2.4": PointSeed("responsibility_exaggeration", "遗产纠纷基本就能解决", "提前把保险配好，家里的遗产纠纷以后基本就能解决。", "inheritance_dispute_claim", compliant_text="保险不应宣传为直接解决遗产纠纷的工具，相关法律问题需按正式程序处理。", keywords=("遗产纠纷",)),
    "2.2.5": PointSeed("responsibility_exaggeration", "能把纠纷解决掉", "就算以后婚姻有变化，这份保单也能帮您把纠纷解决掉。", "marriage_dispute_claim", compliant_text="保险不能承诺解决婚姻纠纷或财产争议，相关法律后果不能夸大。", keywords=("婚姻纠纷",)),
    "2.3.1": PointSeed("responsibility_exaggeration", "一确诊就赔", "这类重疾险一确诊就赔，不用再看别的条件。", "claim_settlement_promise", compliant_text="是否理赔要结合等待期、病种定义和条款约定，不能简单说成确诊即赔。", keywords=("确诊即赔",)),
    "2.3.2": PointSeed("responsibility_exaggeration", "门诊住院还是康复费用，这份保险都能赔", "无论门诊、住院还是康复费用，这份保险基本都能赔。", "claim_scope_exaggeration", compliant_text="门诊、住院和康复费用的赔付范围要按合同责任和报销条件判断。", keywords=("都能赔", "康复费用")),
    "3.1.1": PointSeed("guaranteed_return", "收益是保证的", "这份产品收益是保证的，买进去之后就能安心拿。", "guaranteed_income", compliant_text="利益演示不等于保证收益，实际领取和利益以合同约定及产品机制为准。", keywords=("收益保证",)),
    "3.2.1": PointSeed("guaranteed_return", "保本，还能做到年年返钱", "它相当于保本，还能做到年年返钱。", "fixed_return_claim", compliant_text="不能把保险说成保本或年年返，具体领取安排要看合同条款。", keywords=("保本", "年年返")),
    "3.3.1": PointSeed("guaranteed_return", "收益很稳健，几乎没有波动", "这类产品收益很稳健，几乎没有波动。", "stable_return_claim", compliant_text="产品利益可能波动，不能宣传成没有波动或绝对稳健。", keywords=("稳健", "没有波动")),
    "3.4.1": PointSeed("guaranteed_return", "你做下来也差不多", "我上一个客户买完两年回报就很高，你做下来也差不多。", "case_based_return_implication", compliant_text="客户个案或利益演示不能类比您未来的实际回报，还是要看产品条款。", keywords=("客户案例", "很高回报")),
    "3.5.1": PointSeed("guaranteed_return", "稳赚不赔，几乎零风险", "这就是稳赚不赔的安排，几乎零风险。", "zero_risk_claim", compliant_text="保险产品不能宣传为稳赚不赔或零风险，投保前应了解责任和限制。", keywords=("稳赚不赔", "零风险")),
    "3.6.1": PointSeed("guaranteed_return", "下有保底、上不封顶", "这个方案下有保底、上不封顶，怎么做都划算。", "structured_yield_claim", compliant_text="不能用下有保底、上不封顶替代真实风险提示，具体利益规则要按产品说明理解。", keywords=("下有保底", "上不封顶")),
    "3.8.1": PointSeed("guaranteed_return", "收益明显更高", "和银行存款比，这个收益明显更高。", "relative_yield_claim", compliant_text="保险不能简单和银行存款比收益高低，两类产品功能和风险特征不同。", keywords=("收益更高", "银行")),
    "4.1.1": PointSeed("financial_product_confusion", "别管它是不是保险", "您先别管它是不是保险，本质上就是一笔长期资金工具。", "product_identity_omission", compliant_text="介绍时应明确这是保险产品，并同步说明保障责任、费用和利益边界。", keywords=("别管保险不保险", "长期资金工具")),
    "6.2.1": PointSeed("comparison_violation", "别家的产品理赔慢、服务也差", "别家的产品理赔慢、服务也差，根本没法和我们比。", "competitor_smear", compliant_text="做产品介绍时只客观讲自身特点，不贬低同业或同业产品。", keywords=("理赔慢", "服务差")),
    "6.3.1": PointSeed("comparison_violation", "基金和银行理财都没有这份保险稳", "基金和银行理财都没有这份保险稳。", "financial_product_comparison", compliant_text="不同金融产品定位不同，不应通过片面对比引导客户误解。", keywords=("基金", "银行理财")),
    "6.4.1": PointSeed("comparison_violation", "社保根本解决不了问题", "社保根本解决不了问题，关键还是得靠商业保险。", "social_security_smear", compliant_text="社保和商业保险功能不同，可以互为补充，但不应通过贬低社保销售产品。", keywords=("社保根本解决不了",)),
    "6.5.1": PointSeed("gifts_or_extra_benefits", "返一部分保费给您", "您今天签单的话，我这边可以返一部分保费给您。", "rebate_inducement", compliant_text="投保过程中不得返利、返现或变相返还保费，所有费用应按正规流程缴纳。", keywords=("返保费",)),
    "6.5.2": PointSeed("gifts_or_extra_benefits", "送您一份礼品，再给您一次抽奖机会", "这周投保我送您一份礼品，再给您一次抽奖机会。", "gift_inducement", compliant_text="投保环节不得以礼品、折扣或抽奖诱导成交。", keywords=("送礼品", "抽奖")),
    "7.1.1": PointSeed("absolute_expression", "最好的保险方案", "这是目前市面上最好的保险方案。", "superlative_claim", compliant_text="介绍产品时可以说明适用人群和特点，但不要使用最好、最佳等绝对化用语。", keywords=("最好的",)),
    "7.1.2": PointSeed("absolute_expression", "未来养老一定没问题", "只要买了它，未来养老一定没问题。", "assertive_promise", compliant_text="养老规划需要结合个人情况，不能对未来结果作一定没问题的断言。", keywords=("一定没问题",)),
    "7.1.3": PointSeed("absolute_expression", "独一无二的选择", "这种产品在同类里几乎是独一无二的选择。", "uniqueness_claim", compliant_text="不能用独一无二等说法暗示保险产品具有唯一优势。", keywords=("独一无二",)),
    "7.1.4": PointSeed("absolute_expression", "肯定不会出现金流缺口", "只要现在配置上，未来肯定不会出现金流缺口。", "conditional_absolute_claim", compliant_text="即使做了保险配置，也不能断言未来一定不会出现金流缺口。", keywords=("肯定不会", "资金缺口")),
    "7.2.1": PointSeed("absolute_expression", "再不买就亏大了", "再不买就亏大了，后面想配都来不及。", "fear_pressure_claim", compliant_text="销售时不能用不买就亏、来不及等恐惧式话术催促成交。", keywords=("再不买就亏", "来不及")),
    "7.3.1": PointSeed("absolute_expression", "最后名额", "今天就是最后名额，过了今晚就彻底没有了。", "urgency_claim", compliant_text="限时限量信息必须真实可核验，不能凭空制造最后机会。", keywords=("最后名额",)),
    "7.4.1": PointSeed("absolute_expression", "随时都能取，而且几乎没有任何成本", "这个产品的钱基本随时都能取，而且几乎没有任何成本。", "ambiguous_claim", compliant_text="涉及领取、减保或费用时，应明确条件、限制和可能成本。", keywords=("随时都能取", "没有任何成本")),
    "7.5.1": PointSeed("absolute_expression", "中产家庭都会尽早把这类保险配齐", "真正的中产家庭都会尽早把这类保险配齐。", "improper_term_claim", compliant_text="宣传中应避免使用中产等标签化用语误导客户判断。", keywords=("中产家庭",)),
    "8.2.1": PointSeed("other", "直播间专属福利，点链接现在买最划算", "直播间专属福利，点链接现在买最划算。", "internet_hype", actor_scope="company", compliant_text="互联网渠道宣传也要如实说明产品信息、适用条件和风险提示。", keywords=("直播间专属", "最划算")),
    "8.3.1": PointSeed("responsibility_exaggeration", "随时减保取现，跟活期资金一样灵活", "以后急用钱随时减保取现，跟活期资金一样灵活。", "cash_value_hype", compliant_text="减保或保单贷款需要说明条件、限制和可能影响，不能简单说成像活期一样灵活。", keywords=("随时减保取现", "一样灵活")),
    "9.1.1": PointSeed("other", "国家级保险保障计划", "这是平安官方认证的国家级保险保障计划。", "brand_misuse", actor_scope="company", compliant_text="品牌和机构名称应按授权范围规范使用，不得延伸为国家级背书。", keywords=("国家级", "官方认证")),
    "9.2.1": PointSeed("other", "爆款测评直接拿来做宣传海报", "那篇网上爆款测评我直接拿来做宣传海报就行。", "copyright_violation", actor_scope="company", compliant_text="宣传素材应使用已授权或原创内容，不能直接搬用他人作品。", keywords=("爆款测评", "宣传海报")),
    "9.3.1": PointSeed("other", "行业报告都说这款产品回报最好", "行业报告都说这款产品回报最好，来源我就不展开了。", "unsupported_third_party_data", actor_scope="company", compliant_text="引用第三方数据时应说明来源、时间和适用范围，不能只说报告都这么写。", keywords=("行业报告", "出处")),
    "9.4.1": PointSeed("other", "某位明星都在推荐", "连某位明星都在推荐这份保险，您完全可以放心。", "celebrity_endorsement", actor_scope="company", compliant_text="未经授权，不得借用明星或知名人士名义推荐保险产品。", keywords=("明星", "推荐")),
    "9.5.1": PointSeed("other", "借某银行客户经理的名片一起推广", "我们借某银行客户经理的名片一起推广这款产品。", "third_party_brand_misuse", actor_scope="company", compliant_text="不得借用第三方机构品牌、名片或身份开展保险宣传。", keywords=("银行客户经理", "名片")),
    "9.6.1": PointSeed("other", "奥运冠军同款保障计划", "这是我们的奥运冠军同款保障计划。", "event_ip_misuse", actor_scope="company", compliant_text="不得擅自借用奥运会、世界杯等赛事名称或相关表述营销。", keywords=("奥运", "冠军同款")),
    "10.1.1": PointSeed("gifts_or_extra_benefits", "买这份保险就送居家养老上门照护和定制晚宴服务", "现在投保就送居家养老上门照护和定制晚宴服务。", "service_bundle_inducement", actor_scope="company", compliant_text="健康管理之外的服务宣传应明确边界，不能包装成买保险就送额外服务。", keywords=("送服务", "定制晚宴")),
    "10.2.1": PointSeed("other", "最终解释权归我们所有", "这次活动最终解释权归我们所有，礼品怎么发都由我们定。", "promotion_rule_violation", actor_scope="company", compliant_text="活动规则和奖品发放应事先明确，不得以最终解释权等表述规避责任。", keywords=("最终解释权",)),
    "11.1.1": PointSeed("other", "借着二十大东风", "借着二十大东风，现在正是全民配置保险的最好时机。", "political_association", actor_scope="company", compliant_text="保险宣传不应借用会议名义开展营销，也不应用此制造销售氛围。", keywords=("二十大", "最好时机")),
    "11.2.1": PointSeed("other", "响应国家领导人重要讲话精神推出的重点保障产品", "这是响应国家领导人重要讲话精神推出的重点保障产品。", "leader_association", actor_scope="company", compliant_text="不得使用国家机关或领导人相关表述为具体保险产品背书。", keywords=("国家领导人", "重点产品")),
    "11.3.1": PointSeed("other", "港澳台客户都在抢着配置", "港澳台客户都在抢着配置这款保单。", "sensitive_politics", actor_scope="company", compliant_text="宣传中应避免引入港澳台等敏感政治话题进行引流。", keywords=("港澳台",)),
    "11.3.2": PointSeed("other", "带点辱华争议的话题做海报", "带点辱华争议的话题做海报，更容易出圈。", "sensitive_politics", actor_scope="company", compliant_text="不得借辱华或其他敏感政治内容吸引流量或制造噱头。", keywords=("辱华",)),
    "11.3.3": PointSeed("other", "用文革式口号包装", "用文革式口号包装我们的增员海报，会更有冲击力。", "sensitive_politics", actor_scope="company", compliant_text="不得使用文革等敏感政治内容做产品或增员宣传。", keywords=("文革",)),
    "11.4.1": PointSeed("other", "放国旗国徽，更能显得这款保险有权威性", "海报上放国旗国徽，更能显得这款保险有权威性。", "national_symbol_misuse", actor_scope="company", compliant_text="不得在保险营销中违规使用国旗、国徽、国歌等元素。", keywords=("国旗国徽", "权威性")),
    "11.6.1": PointSeed("other", "趁着这次公共卫生事件关注高", "趁着这次公共卫生事件关注高，正好推我们的医疗险。", "hot_event_marketing", actor_scope="company", compliant_text="不得借公共卫生事件、灾害事故等社会热点进行营销炒作。", keywords=("公共卫生事件",)),
    "12.1.1": PointSeed("agent_title_violation", "直接写家庭理财规划师", "对外别写保险代理人，直接写家庭理财规划师。", "title_misuse", compliant_text="对外展业应如实使用保险代理人等合规称谓，不得自行改称。", keywords=("理财规划师", "代理人")),
    "12.2.1": PointSeed("agent_title_violation", "轻松月入过万，收入上不封顶", "加入我们轻松月入过万，收入上不封顶。", "income_recruitment", compliant_text="招聘宣传不得用月入过万、上不封顶等收入承诺诱导加入。", keywords=("月入过万", "上不封顶")),
    "12.3.1": PointSeed("agent_title_violation", "什么金融产品都可以顺带一起讲", "拿到工号以后，什么金融产品都可以顺带一起讲。", "scope_confusion", compliant_text="代理人的展业范围和资格应按监管要求如实说明，不得超范围表述。", keywords=("什么金融产品都能一起讲",)),
    "12.4.1": PointSeed("agent_title_violation", "这不是招募代理人，就是在招企业合伙人", "这不是招募代理人，就是在招企业合伙人。", "recruitment_disguise", compliant_text="招募代理人时应如实披露合作关系和职业性质，不能包装成其他职业。", keywords=("企业合伙人", "招募代理人")),
    "13.1.1": PointSeed("regulatory_misinterpretation", "国家鼓励推广的保险产品", "这是国家鼓励推广的保险产品，早点配最放心。", "national_endorsement", actor_scope="company", compliant_text="政策背景不等于国家为具体保险产品背书，宣传时应避免这种引导。", keywords=("国家鼓励",)),
    "13.2.1": PointSeed("regulatory_misinterpretation", "监管都备案了，就说明这款产品绝对安全", "既然监管都备案了，就说明这款产品绝对安全。", "regulator_endorsement", actor_scope="company", compliant_text="备案或审批程序不代表监管机关为产品收益、风险和销售口径背书。", keywords=("监管备案", "绝对安全")),
    "13.3.1": PointSeed("regulatory_misinterpretation", "某位领导都强调要重视这类保障，现在买最合适", "某位领导都强调要重视这类保障，现在买最合适。", "leadership_endorsement", actor_scope="company", compliant_text="不得借用领导人讲话内容对具体保险产品作引导性宣传。", keywords=("领导", "最合适")),
    "13.4.1": PointSeed("regulatory_misinterpretation", "保险公司不会倒闭，所以这份保单绝对安全", "保险公司不会倒闭，所以这份保单绝对安全。", "policy_endorsement", actor_scope="company", compliant_text="不得以保险公司稳定性替代对产品风险、责任和限制的如实说明。", keywords=("不会倒闭", "绝对安全")),
    "14.1.1": PointSeed("other", "开过光，还能祈福保平安", "这份保单开过光，还能祈福保平安，给家人更安心。", "religion_inducement", actor_scope="company", compliant_text="宣传内容不得涉及开光、祈福、菩萨、财神等宗教诱导表述。", keywords=("开过光", "祈福")),
}


COMPLIANT_MITIGATIONS = {
    "financial_product_confusion": ["clarified_insurance_nature", "contract_disclaimer"],
    "guaranteed_return": ["uncertainty_disclosure", "contract_disclaimer"],
    "gifts_or_extra_benefits": ["no_extra_benefits", "sales_boundary_clarified"],
    "responsibility_exaggeration": ["coverage_boundary_clarified", "contract_disclaimer"],
    "absolute_expression": ["neutral_wording", "no_absolute_claim"],
    "comparison_violation": ["objective_description", "no_competitor_smear"],
    "agent_title_violation": ["identity_disclosed", "scope_disclosed"],
    "regulatory_misinterpretation": ["no_regulatory_endorsement", "objective_policy_reference"],
    "other": ["neutral_wording"],
}


COMPLIANT_EXCEPTION_MARKERS = ("不得", "不能", "不应", "不可", "避免", "不是", "不能简单", "不等于")


def extract_point_id(point_text: str) -> str:
    return point_text.split(" ", 1)[0].strip()


def build_violation_text(seed: PointSeed) -> str:
    return seed.violation_text


def build_compliant_text(seed: PointSeed) -> str:
    return seed.compliant_text


def infer_exception_triggered(text: str) -> bool:
    return any(marker in text for marker in COMPLIANT_EXCEPTION_MARKERS)


def make_sample_id(prefix: str, point_id: str) -> str:
    return f"{prefix}_{point_id.replace('.', '_')}"


def build_records(excel_path: Path) -> list[CaseRecord]:
    wb = load_workbook(excel_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    records: list[CaseRecord] = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        level1 = str(row[0] or "").strip()
        level2 = str(row[1] or "").strip()
        point_text = str(row[2] or "").strip()
        if not point_text:
            continue

        point_id = extract_point_id(point_text)
        if point_id not in POINT_SEEDS:
            raise KeyError(f"未为审查点配置种子模板: {point_id} / {point_text}")

        seed = POINT_SEEDS[point_id]
        keywords = list(seed.keywords)

        violation_record = CaseRecord(
            sample_id=make_sample_id("审查点违规", point_id),
            label="violation",
            text=build_violation_text(seed),
            source_sheet="三级审查点",
            source_row=row_idx,
            source_column="三级违规点",
            reason_raw=point_text,
            expected_categories=[seed.category],
            expected_text_slices=[seed.violation_phrase],
            expected_keywords=keywords,
            tags=["audit_point_seed", "synthetic_smoke", "point_coverage"],
            expected_audit_point_id=point_id,
            expected_exception_triggered=False,
            expected_actor_scope=seed.actor_scope,
            expected_claim_type=seed.claim_type,
            is_hard_negative=False,
        )
        records.append(violation_record)

        compliant_record = CaseRecord(
            sample_id=make_sample_id("审查点合规", point_id),
            label="compliant",
            text=build_compliant_text(seed),
            source_sheet="三级审查点",
            source_row=row_idx,
            source_column="三级违规点",
            reason_raw=f"{point_text} 的合规对照表达",
            expected_categories=[],
            expected_text_slices=[],
            expected_keywords=keywords,
            tags=["audit_point_seed", "synthetic_smoke", "point_coverage", "hard_negative"],
            expected_audit_point_id=point_id,
            expected_exception_triggered=infer_exception_triggered(seed.compliant_text),
            expected_mitigating_factors=COMPLIANT_MITIGATIONS.get(seed.category, ["neutral_wording"]),
            expected_actor_scope=seed.actor_scope,
            expected_claim_type=seed.claim_type,
            is_hard_negative=True,
        )
        records.append(compliant_record)

    return records


def write_jsonl(path: Path, records: list[CaseRecord]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        for record in records:
            fh.write(json.dumps(record.to_dict(), ensure_ascii=False) + "\n")


def write_summary(path: Path, records: list[CaseRecord], point_count: int) -> None:
    by_label = Counter(record.label for record in records)
    by_category = Counter()
    hard_negatives = 0
    for record in records:
        if record.is_hard_negative:
            hard_negatives += 1
        for category in record.expected_categories:
            by_category[category] += 1

    lines = [
        "# 审查点覆盖型 Smoke 构造摘要",
        "",
        "- 数据来源：`plan/smoke数据构造种子集.xlsx`",
        f"- 三级审查点数：`{point_count}`",
        f"- 输出样本数：`{len(records)}`",
        f"- 其中违规样本：`{by_label.get('violation', 0)}`",
        f"- 其中合规对照样本：`{by_label.get('compliant', 0)}`",
        f"- hard negative 样本：`{hard_negatives}`",
        "",
        "## 设计原则",
        "- 每个三级审查点生成 1 条违规样本 + 1 条合规对照样本",
        "- 数据为覆盖型 synthetic smoke，用于验证规则覆盖与边界判断，不替代真实语料评测",
        "- 所有样本均带 `expected_audit_point_id`，便于后续做审查点级命中分析",
        "",
        "## 违规类别覆盖",
    ]

    for category, count in sorted(by_category.items()):
        lines.append(f"- `{category}`: `{count}`")

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="根据审查点种子表构造 smoke 数据")
    parser.add_argument(
        "--input",
        default="plan/smoke数据构造种子集.xlsx",
        help="审查点种子 Excel 路径",
    )
    parser.add_argument(
        "--output",
        default="benchmark/datasets/smoke/case_eval_smoke_audit_points_20260318.jsonl",
        help="输出 JSONL 路径",
    )
    parser.add_argument(
        "--summary",
        default="benchmark/reports/smoke_build_audit_points_20260318.md",
        help="摘要报告路径",
    )
    args = parser.parse_args()

    input_path = Path(args.input)
    records = build_records(input_path)
    point_count = len(records) // 2

    write_jsonl(Path(args.output), records)
    write_summary(Path(args.summary), records, point_count)

    print(
        json.dumps(
            {
                "input": args.input,
                "points": point_count,
                "samples": len(records),
                "output": args.output,
                "summary": args.summary,
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
