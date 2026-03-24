"""
三级审查点目录与规则富化
========================
基于 `plan/smoke数据构造种子集.xlsx` 构建审查点目录，
并在运行时为 RuleCard 自动补齐：
  - audit_point_id / audit_point_name
  - primary_category / secondary_category
  - actor_scope / claim_type
  - route_hint / complexity_level

目标：
1. 让“规则 -> 审查点 -> 输出类别”形成可显式追踪的闭环
2. 为 benchmark 提供更稳定的审查点级评估基础
3. 优先收口“金融用语混淆”“绝对化夸大表述”两类重点能力
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from . import config
from .schemas import RuleCard


@dataclass(frozen=True)
class AuditPointDefinition:
    audit_point_id: str
    audit_point_name: str
    level1_name: str
    level2_name: str
    primary_category: str
    secondary_category: str
    category_group: str | None
    actor_scope: str | None
    claim_type: str | None
    route_hint: str
    complexity_level: str
    complex_skill_type: str | None
    keywords: tuple[str, ...]


@dataclass(frozen=True)
class AuditPointPreset:
    primary_category: str
    secondary_category: str
    category_group: str | None = None
    actor_scope: str | None = None
    claim_type: str | None = None
    route_hint: str = "neutral"
    complexity_level: str = "simple"
    complex_skill_type: str | None = None
    keywords: tuple[str, ...] = ()


_CATEGORY_GROUP_TO_PRIMARY_CATEGORY = {
    "financial_confusion": "financial_product_confusion",
    "guaranteed_return": "guaranteed_return",
    "gifts_benefits": "gifts_or_extra_benefits",
    "responsibility_exaggeration": "responsibility_exaggeration",
    "absolute_expression": "absolute_expression",
    "regulatory_misinterpretation": "regulatory_misinterpretation",
    "surrender_guidance": "surrender_guidance",
    "agent_title_violation": "agent_title_violation",
    "comparison_violation": "comparison_violation",
    "other": "other",
}


_AUDIT_POINT_PRESETS: dict[str, AuditPointPreset] = {
    "1.1.1": AuditPointPreset("financial_product_confusion", "deposit_payment_confusion", "financial_confusion", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="financial_confusion", keywords=("存", "存入", "存款", "缴费", "交费")),
    "1.1.2": AuditPointPreset("financial_product_confusion", "principal_interest_confusion", "financial_confusion", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="financial_confusion", keywords=("本金", "利息", "复利", "计息")),
    "1.1.3": AuditPointPreset("financial_product_confusion", "savings_account_confusion", "financial_confusion", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="financial_confusion", keywords=("账户", "储蓄账户", "万能账户")),
    "1.2.1": AuditPointPreset("financial_product_confusion", "bank_fund_analogy", "financial_confusion", claim_type="comparison_advantage", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="financial_confusion", keywords=("银行理财", "基金", "类比", "银行", "理财")),
    "2.1.1": AuditPointPreset("responsibility_exaggeration", "coverage_absolute_claim", "responsibility_exaggeration", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="comparison_or_absolute", keywords=("保障范围", "理赔范围", "全赔", "全额赔", "什么都赔")),
    "2.2.1": AuditPointPreset("responsibility_exaggeration", "debt_planning_claim", "responsibility_exaggeration", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="tax_or_law_misinterpretation", keywords=("债务", "避债", "债务规划")),
    "2.2.2": AuditPointPreset("responsibility_exaggeration", "tax_planning_claim", "responsibility_exaggeration", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="tax_or_law_misinterpretation", keywords=("税务", "税收", "避税", "免税")),
    "2.2.3": AuditPointPreset("responsibility_exaggeration", "universal_coverage_claim", "responsibility_exaggeration", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", keywords=("保所有", "什么都保", "全部都保")),
    "2.2.4": AuditPointPreset("responsibility_exaggeration", "inheritance_dispute_resolution", "responsibility_exaggeration", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="transfer_or_inheritance", keywords=("遗产纠纷", "财富传承", "财产分割", "遗产")),
    "2.2.5": AuditPointPreset("responsibility_exaggeration", "marriage_dispute_resolution", "responsibility_exaggeration", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="transfer_or_inheritance", keywords=("婚姻纠纷", "离婚", "财产分割", "法律纠纷")),
    "2.3.1": AuditPointPreset("responsibility_exaggeration", "claim_condition_simplification", "responsibility_exaggeration", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", keywords=("确诊即赔", "理赔条件", "简化理赔", "马上赔")),
    "2.3.2": AuditPointPreset("responsibility_exaggeration", "claim_scope_exaggeration", "responsibility_exaggeration", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", keywords=("理赔范围", "都能赔", "都可以赔", "赔付范围")),
    "3.1.1": AuditPointPreset("guaranteed_return", "deterministic_return_claim", "guaranteed_return", claim_type="income_promise", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="guaranteed_return", keywords=("保证收益", "确定性收益", "锁定收益", "回报")),
    "3.2.1": AuditPointPreset("guaranteed_return", "fixed_return_claim", "guaranteed_return", claim_type="income_promise", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="guaranteed_return", keywords=("保本", "年年返", "固定收益")),
    "3.3.1": AuditPointPreset("guaranteed_return", "stable_return_implication", "guaranteed_return", claim_type="income_promise", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="guaranteed_return", keywords=("稳健", "无波动", "稳定收益", "稳稳的")),
    "3.4.1": AuditPointPreset("guaranteed_return", "case_based_return_analogy", "guaranteed_return", claim_type="historical_performance", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="guaranteed_return", keywords=("客户案例", "案例收益", "未来回报", "类比回报", "未来经济收益")),
    "3.5.1": AuditPointPreset("guaranteed_return", "rhetorical_guarantee", "guaranteed_return", claim_type="income_promise", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="guaranteed_return", keywords=("稳赚不赔", "零风险", "稳赚", "零波动")),
    "3.6.1": AuditPointPreset("guaranteed_return", "template_return_structure", "guaranteed_return", claim_type="income_promise", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="guaranteed_return", keywords=("下有保底", "上不封顶", "保底", "无上限")),
    "3.8.1": AuditPointPreset("guaranteed_return", "relative_bank_return_claim", "guaranteed_return", claim_type="comparison_advantage", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="guaranteed_return", keywords=("比银行收益更高", "比银行", "收益更高", "比存银行", "远高于")),
    "6.2.1": AuditPointPreset("comparison_violation", "peer_product_derogation", "comparison_violation", claim_type="comparison_advantage", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="comparison_or_absolute", keywords=("贬低同业", "同业产品", "他司产品", "诋毁")),
    "6.3.1": AuditPointPreset("comparison_violation", "one_sided_financial_comparison", "comparison_violation", claim_type="comparison_advantage", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="financial_confusion", keywords=("片面对比", "金融产品", "银行", "基金", "理财")),
    "6.5.1": AuditPointPreset("gifts_or_extra_benefits", "rebate_or_premium_return", "gifts_benefits", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="gifts_or_extra_benefits", keywords=("回扣", "返还保费", "返利", "返佣")),
    "6.5.2": AuditPointPreset("gifts_or_extra_benefits", "gift_discount_lottery_inducement", "gifts_benefits", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="gifts_or_extra_benefits", keywords=("礼品", "折扣", "抽奖", "惊喜", "温馨服务")),
    "7.1.1": AuditPointPreset("absolute_expression", "superlative_claim", "absolute_expression", claim_type="ranking_claim", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="comparison_or_absolute", keywords=("最好", "最佳", "第一", "最优", "no.1")),
    "7.1.2": AuditPointPreset("absolute_expression", "assertive_promise", "absolute_expression", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="comparison_or_absolute", keywords=("一定", "肯定", "绝对", "必然")),
    "7.1.3": AuditPointPreset("absolute_expression", "uniqueness_claim", "absolute_expression", claim_type="ranking_claim", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="comparison_or_absolute", keywords=("唯一", "独一无二", "不可替代")),
    "7.1.4": AuditPointPreset("absolute_expression", "conditional_absolute_claim", "absolute_expression", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="comparison_or_absolute", keywords=("只要", "一定", "肯定不会", "不会出现金流缺口")),
    "7.2.1": AuditPointPreset("absolute_expression", "fear_pressure_claim", "absolute_expression", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="comparison_or_absolute", keywords=("不买就亏", "再不买就晚", "来不及")),
    "7.3.1": AuditPointPreset("absolute_expression", "urgency_claim", "absolute_expression", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="comparison_or_absolute", keywords=("限时限量", "最后名额", "过了今晚", "马上抢", "限时", "限量", "先到先得", "欲购从速")),
    "7.5.1": AuditPointPreset("absolute_expression", "improper_term_claim", "absolute_expression", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="comparison_or_absolute", keywords=("中产", "精英家庭", "高净值")),
    "8.2.1": AuditPointPreset("other", "internet_hype_claim", "other", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", keywords=("爆款", "直播间", "朋友圈", "扫码", "平台营销")),
    "9.1.1": AuditPointPreset("other", "brand_scope_misuse", "other", claim_type="misleading_statement", route_hint="prefer_base", complexity_level="simple", keywords=("品牌", "禁用词", "机构名称", "公司名称")),
    "9.2.1": AuditPointPreset("other", "copyright_or_content_violation", "other", claim_type="misleading_statement", route_hint="prefer_base", complexity_level="simple", keywords=("未经授权", "版权", "违规内容", "logo")),
    "9.4.1": AuditPointPreset("other", "celebrity_endorsement_infringement", "other", claim_type="misleading_statement", route_hint="prefer_base", complexity_level="simple", keywords=("明星", "知名人士", "代言", "名义宣传")),
    "9.5.1": AuditPointPreset("other", "third_party_business_card_misuse", "other", claim_type="misleading_statement", route_hint="prefer_base", complexity_level="simple", keywords=("第三方名片", "名片", "借用名片", "品牌名片")),
    "9.6.1": AuditPointPreset("other", "sports_event_brand_misuse", "other", claim_type="misleading_statement", route_hint="prefer_base", complexity_level="simple", keywords=("奥运", "世界杯", "奥林匹克", "亚运会")),
    "10.1.1": AuditPointPreset("gifts_or_extra_benefits", "service_beyond_health_management", "gifts_benefits", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="gifts_or_extra_benefits", keywords=("买保险送服务", "超出健康管理", "送服务")),
    "10.2.1": AuditPointPreset("other", "activity_rule_violation", "other", claim_type="misleading_statement", route_hint="prefer_base", complexity_level="simple", keywords=("解释权", "奖品超限", "活动规则", "专属奖", "奖品")),
    "11.1.1": AuditPointPreset("other", "party_or_conference_reference", "other", claim_type="misleading_statement", route_hint="prefer_base", complexity_level="simple", keywords=("十九大", "二十大", "建党百年")),
    "11.2.1": AuditPointPreset("other", "leader_or_agency_name_use", "other", claim_type="misleading_statement", route_hint="prefer_base", complexity_level="simple", keywords=("国家领导", "国家机关", "工作人员", "形象")),
    "11.3.1": AuditPointPreset("other", "sensitive_political_hk_tw", "other", claim_type="misleading_statement", route_hint="prefer_base", complexity_level="simple", keywords=("港澳台", "台湾")),
    "11.3.2": AuditPointPreset("other", "sensitive_political_insulting_china", "other", claim_type="misleading_statement", route_hint="prefer_base", complexity_level="simple", keywords=("辱华",)),
    "11.3.3": AuditPointPreset("other", "sensitive_political_cultural_revolution", "other", claim_type="misleading_statement", route_hint="prefer_base", complexity_level="simple", keywords=("文革",)),
    "11.4.1": AuditPointPreset("other", "national_flag_emblem_anthem", "other", claim_type="misleading_statement", route_hint="prefer_base", complexity_level="simple", keywords=("国旗", "国徽", "国歌")),
    "11.6.1": AuditPointPreset("other", "hot_event_leveraging", "other", claim_type="misleading_statement", route_hint="prefer_base", complexity_level="simple", keywords=("自然灾害", "事故灾难", "公共卫生事件", "热点事件")),
    "12.1.1": AuditPointPreset("agent_title_violation", "misleading_position_title", "agent_title_violation", actor_scope="agent", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="agent_title_or_recruitment", keywords=("顾问", "总监", "经理", "职位", "金融顾问", "财富顾问", "规划师", "财富管理师", "产品专家", "医养顾问")),
    "12.2.1": AuditPointPreset("agent_title_violation", "income_exaggeration_recruitment", "agent_title_violation", actor_scope="agent", claim_type="income_promise", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="agent_title_or_recruitment", keywords=("月入过万", "无上限", "高薪", "收入", "月薪", "月收入")),
    "12.3.1": AuditPointPreset("agent_title_violation", "scope_of_practice_confusion", "agent_title_violation", actor_scope="agent", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="agent_title_or_recruitment", keywords=("展业范围", "代理人", "职责范围")),
    "12.4.1": AuditPointPreset("agent_title_violation", "recruitment_nature_confusion", "agent_title_violation", actor_scope="agent", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="agent_title_or_recruitment", keywords=("招募", "招聘", "职业行为", "合伙人", "应聘", "创业")),
    "13.1.1": AuditPointPreset("regulatory_misinterpretation", "national_endorsement_claim", "regulatory_misinterpretation", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="national_or_regulatory_endorsement", keywords=("国家鼓励", "国家支持")),
    "13.2.1": AuditPointPreset("regulatory_misinterpretation", "regulatory_approval_endorsement", "regulatory_misinterpretation", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="national_or_regulatory_endorsement", keywords=("监管机构", "备案程序", "审核通过", "监管审批")),
    "13.3.1": AuditPointPreset("regulatory_misinterpretation", "leader_endorsement_claim", "regulatory_misinterpretation", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="national_or_regulatory_endorsement", keywords=("领导人讲话", "政府机构领导", "领导背书", "副主席", "国家机关工作人员")),
    "13.4.1": AuditPointPreset("regulatory_misinterpretation", "policy_backing_claim", "regulatory_misinterpretation", claim_type="misleading_statement", route_hint="prefer_skill", complexity_level="complex", complex_skill_type="national_or_regulatory_endorsement", keywords=("不会倒闭", "政策支持", "国家兜底", "不会破产", "不会解散")),
    "14.1.1": AuditPointPreset("other", "religious_inducement", "other", claim_type="misleading_statement", route_hint="prefer_base", complexity_level="simple", keywords=("祈福", "开光", "菩萨", "财神", "法师")),
}


def _normalize_text(text: str) -> str:
    return str(text or "").strip().lower()


def _extract_keywords(*texts: str) -> tuple[str, ...]:
    candidates: list[str] = []
    for text in texts:
        normalized = str(text or "").strip()
        if not normalized:
            continue
        quoted = re.findall(r"[\"“](.*?)[\"”]", normalized)
        for item in quoted:
            candidates.extend(re.split(r"[、/／，,；; ]+", item))
        candidates.extend(re.findall(r"[A-Za-z0-9\.]+|[\u4e00-\u9fff]{2,}", normalized))
    cleaned: list[str] = []
    seen: set[str] = set()
    for item in candidates:
        term = item.strip("：:（）()“”\"'，,。 ")
        if len(term) < 2 or term in seen:
            continue
        seen.add(term)
        cleaned.append(term)
    return tuple(cleaned)


def _default_primary_category(audit_point_id: str) -> str:
    if audit_point_id.startswith("1."):
        return "financial_product_confusion"
    if audit_point_id.startswith("2."):
        return "responsibility_exaggeration"
    if audit_point_id.startswith("3."):
        return "guaranteed_return"
    if audit_point_id.startswith("6.2") or audit_point_id.startswith("6.3"):
        return "comparison_violation"
    if audit_point_id.startswith("6.5") or audit_point_id.startswith("10.1"):
        return "gifts_or_extra_benefits"
    if audit_point_id.startswith("7."):
        return "absolute_expression"
    if audit_point_id.startswith("12."):
        return "agent_title_violation"
    if audit_point_id.startswith("13."):
        return "regulatory_misinterpretation"
    return "other"


def _default_category_group(primary_category: str) -> str | None:
    reverse = {
        "financial_product_confusion": "financial_confusion",
        "guaranteed_return": "guaranteed_return",
        "gifts_or_extra_benefits": "gifts_benefits",
        "responsibility_exaggeration": "responsibility_exaggeration",
        "absolute_expression": "absolute_expression",
        "regulatory_misinterpretation": "regulatory_misinterpretation",
        "surrender_guidance": "surrender_guidance",
        "agent_title_violation": "agent_title_violation",
        "comparison_violation": "comparison_violation",
        "other": "other",
    }
    return reverse.get(primary_category, "other")


@lru_cache(maxsize=1)
def load_audit_point_catalog() -> dict[str, AuditPointDefinition]:
    workbook_path = Path(config.AUDIT_POINT_WORKBOOK_PATH)
    wb = load_workbook(workbook_path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    catalog: dict[str, AuditPointDefinition] = {}
    for level1_name, level2_name, raw_level3_name in ws.iter_rows(min_row=2, values_only=True):
        match = re.match(r"^(\d+(?:\.\d+){2})\s*(.*)$", str(raw_level3_name or ""))
        if not match:
            continue

        audit_point_id, audit_point_name = match.group(1), match.group(2)
        preset = _AUDIT_POINT_PRESETS.get(audit_point_id)
        primary_category = preset.primary_category if preset else _default_primary_category(audit_point_id)
        category_group = preset.category_group if preset else _default_category_group(primary_category)
        secondary_category = (
            preset.secondary_category
            if preset
            else re.sub(r"[^a-z0-9_]+", "_", _normalize_text(audit_point_name)).strip("_") or "other"
        )
        keywords = tuple(dict.fromkeys(
            (
                *(preset.keywords if preset else ()),
                *_extract_keywords(str(level2_name), audit_point_name),
            )
        ))

        catalog[audit_point_id] = AuditPointDefinition(
            audit_point_id=audit_point_id,
            audit_point_name=audit_point_name,
            level1_name=str(level1_name or ""),
            level2_name=str(level2_name or ""),
            primary_category=primary_category,
            secondary_category=secondary_category,
            category_group=category_group,
            actor_scope=preset.actor_scope if preset else None,
            claim_type=preset.claim_type if preset else None,
            route_hint=preset.route_hint if preset else "neutral",
            complexity_level=preset.complexity_level if preset else "simple",
            complex_skill_type=preset.complex_skill_type if preset else None,
            keywords=keywords,
        )

    return catalog


def _build_rule_text(rule_card: RuleCard) -> str:
    return " ".join(
        part
        for part in [
            rule_card.rule_name,
            rule_card.violation_definition,
            " ".join(rule_card.keywords),
            " ".join(rule_card.violation_terms),
            " ".join(rule_card.condition_terms),
            " ".join(rule_card.exclusion_terms),
            rule_card.compliant_basis,
            rule_card.violation_basis,
            rule_card.compliant_case,
            rule_card.violation_case,
        ]
        if part
    ).lower()


def _iter_category_candidates(
    catalog: dict[str, AuditPointDefinition],
    rule_card: RuleCard,
) -> Iterable[AuditPointDefinition]:
    target_primary = (
        rule_card.primary_category
        or _CATEGORY_GROUP_TO_PRIMARY_CATEGORY.get(rule_card.category_group or "", "")
    )
    if target_primary:
        matched = [point for point in catalog.values() if point.primary_category == target_primary]
        if matched:
            return matched
    return catalog.values()


def infer_audit_point_for_rule(rule_card: RuleCard) -> AuditPointDefinition | None:
    """根据规则内容为 RuleCard 推断最可能的三级审查点。"""

    if rule_card.audit_point_id:
        return load_audit_point_catalog().get(rule_card.audit_point_id)

    catalog = load_audit_point_catalog()
    rule_text = _build_rule_text(rule_card)
    best_point: AuditPointDefinition | None = None
    best_score = 0

    for point in _iter_category_candidates(catalog, rule_card):
        score = 0
        if point.category_group and rule_card.category_group == point.category_group:
            score += 4
        if point.primary_category and rule_card.primary_category == point.primary_category:
            score += 4
        if point.actor_scope and rule_card.actor_scope == point.actor_scope:
            score += 2
        if point.claim_type and rule_card.claim_type == point.claim_type:
            score += 2

        for keyword in point.keywords:
            normalized_keyword = _normalize_text(keyword)
            if not normalized_keyword:
                continue
            if normalized_keyword in rule_text:
                score += 3 if len(normalized_keyword) >= 3 else 1

        if score > best_score:
            best_score = score
            best_point = point

    if best_point and best_score >= 3:
        return best_point
    return None


def enrich_rule_card(rule_card: RuleCard) -> RuleCard:
    """基于三级审查点目录对 RuleCard 做运行时富化。"""

    point = infer_audit_point_for_rule(rule_card)
    updates: dict[str, object] = {}

    if point:
        if not rule_card.audit_point_id:
            updates["audit_point_id"] = point.audit_point_id
        if not rule_card.audit_point_name:
            updates["audit_point_name"] = point.audit_point_name
        if not rule_card.primary_category:
            updates["primary_category"] = point.primary_category
        if not rule_card.secondary_category:
            updates["secondary_category"] = point.secondary_category
        if not rule_card.category_group and point.category_group:
            updates["category_group"] = point.category_group
        if not rule_card.actor_scope and point.actor_scope:
            updates["actor_scope"] = point.actor_scope
        if not rule_card.claim_type and point.claim_type:
            updates["claim_type"] = point.claim_type
        if (
            rule_card.route_strategy == "auto"
            and rule_card.route_hint in {None, "neutral"}
            and point.route_hint != "neutral"
        ):
            updates["route_hint"] = point.route_hint
        if rule_card.complexity_level == "simple" and point.complexity_level == "complex":
            updates["complexity_level"] = "complex"

    if not updates and not rule_card.primary_category and rule_card.category_group:
        updates["primary_category"] = _CATEGORY_GROUP_TO_PRIMARY_CATEGORY.get(rule_card.category_group, "other")

    return rule_card.model_copy(update=updates) if updates else rule_card


def build_audit_point_coverage(rule_cards: dict[str, RuleCard]) -> dict[str, list[str]]:
    """构建“审查点 -> 规则 ID”覆盖矩阵。"""

    coverage = {point_id: [] for point_id in load_audit_point_catalog()}
    for rule_card in rule_cards.values():
        enriched = enrich_rule_card(rule_card)
        if enriched.audit_point_id:
            coverage.setdefault(enriched.audit_point_id, []).append(enriched.rule_id)
    return coverage
